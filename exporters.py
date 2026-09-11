import os
import json
import re
import docx
from io import BytesIO
from backend.database import (
    get_or_create_default_novel, list_chapters, list_story_bible_entries,
    get_canon_history, get_project_settings, create_chapter
)

def export_manuscript_text(novel_id: str, format_type: str = "txt") -> str:
    chapters = list_chapters(novel_id, include_archived=False)
    out = []
    
    for ch in chapters:
        title_line = f"Chapter {ch['chapter_number']}: {ch['title']}"
        if format_type == "markdown":
            out.append(f"# {title_line}\n")
            if ch.get("pov"):
                out.append(f"*POV: {ch['pov']}*\n")
        else:
            out.append(title_line + "\n" + ("=" * len(title_line)) + "\n")
            if ch.get("pov"):
                out.append(f"POV: {ch['pov']}\n")
        
        # Strip HTML for clean prose
        clean_content = ch["content"] or ""
        # Convert <p> to double newlines
        clean_content = re.sub(r'</p>', '\n\n', clean_content)
        clean_content = re.sub(r'<br\s*/?>', '\n', clean_content)
        clean_content = re.sub(r'<[^>]+>', '', clean_content)
        clean_content = re.sub(r'\n{3,}', '\n\n', clean_content).strip()
        
        out.append(clean_content + "\n\n")
        if format_type == "markdown":
            out.append("---\n\n")
        else:
            out.append("\n* * *\n\n")
            
    return "".join(out)

def export_manuscript_docx(novel_id: str) -> BytesIO:
    doc = docx.Document()
    novel = get_or_create_default_novel()
    
    # Title Page
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(novel["title"])
    title_run.font.size = docx.shared.Pt(24)
    title_run.font.bold = True
    
    if novel.get("description"):
        desc_p = doc.add_paragraph()
        desc_run = desc_p.add_run(novel["description"])
        desc_run.font.italic = True
    
    doc.add_page_break()
    
    chapters = list_chapters(novel_id, include_archived=False)
    for ch in chapters:
        h = doc.add_heading(f"Chapter {ch['chapter_number']}: {ch['title']}", level=1)
        if ch.get("pov"):
            p_pov = doc.add_paragraph()
            r_pov = p_pov.add_run(f"POV: {ch['pov']}")
            r_pov.font.italic = True
            
        clean_content = ch["content"] or ""
        # Split by paragraphs
        paras = re.findall(r'<p[^>]*>(.*?)</p>', clean_content, re.DOTALL | re.IGNORECASE)
        if not paras:
            paras = clean_content.splitlines()
            
        for p_text in paras:
            p_clean = re.sub(r'<[^>]+>', '', p_text).strip()
            if p_clean:
                doc.add_paragraph(p_clean)
                
        doc.add_page_break()
        
    f = BytesIO()
    doc.save(f)
    f.seek(0)
    return f

def export_full_project_json(novel_id: str) -> dict:
    novel = get_or_create_default_novel()
    chapters = list_chapters(novel_id, include_archived=True)
    story_bible = list_story_bible_entries(novel_id)
    canon_history = get_canon_history(novel_id, limit=500)
    settings = get_project_settings(novel_id)
    
    return {
        "export_version": "1.0",
        "app": "Ascendancy Writer",
        "novel": novel,
        "chapters": chapters,
        "story_bible": story_bible,
        "canon_history": canon_history,
        "settings": settings
    }

def import_chapter_file(novel_id: str, filename: str, file_bytes: bytes) -> dict:
    fname = filename.lower()
    content_text = ""
    title = os.path.splitext(filename)[0]
    
    if fname.endswith(".docx"):
        doc = docx.Document(BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content_text = "".join([f"<p>{p}</p>" for p in paragraphs])
    else:
        raw = file_bytes.decode("utf-8", errors="replace")
        paragraphs = [p.strip() for p in raw.split("\n\n") if p.strip()]
        content_text = "".join([f"<p>{p}</p>" for p in paragraphs])
        
    return create_chapter(
        novel_id=novel_id,
        title=title,
        content=content_text,
        status="Draft",
        notes=f"Imported from {filename}"
    )
